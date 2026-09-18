from io import BytesIO
from zipfile import BadZipFile, ZipFile

from django.contrib.auth.models import User
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.exceptions import ValidationError as DRFValidationError

from ops import models, serializers
from ops.tenancy import organization_scope_ids

from .license_access import license_user_limit_denial
from .phone_identity import normalize_mobile_phone, phone_lookup_hash


MAX_IMPORT_FILE_SIZE = 5 * 1024 * 1024
MAX_IMPORT_ARCHIVE_SIZE = 25 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
TEMPLATE_HEADERS = (
    '用户名', '姓名', '所属部门路径', '邮箱', '手机号',
    '身份', '职位', '启用状态', 'OTP策略', '初始密码',
)
ACTIVE_VALUES = {
    '是': True, '启用': True, '允许': True, 'true': True, '1': True,
    '否': False, '禁用': False, '不允许': False, 'false': False, '0': False,
}
OTP_VALUES = {
    '': 'inherit', '跟随平台': 'inherit', 'inherit': 'inherit',
    '强制启用': 'required', 'required': 'required',
    '免于认证': 'exempt', 'exempt': 'exempt',
}


class UserImportError(ValueError):
    """表示用户 Excel 文件结构、内容或最终导入状态不符合要求。"""

    def __init__(self, message, result=None):
        """保存用户友好错误以及可选的逐行预检查结果。"""
        super().__init__(message)
        self.result = result


def _cell_text(value):
    """把 Excel 单元格值转换为稳定文本，避免整数被显示为小数。"""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized_department_path(value):
    """统一部门路径两侧空白和分隔符，供模板与导入校验一致匹配。"""
    return '/'.join(part.strip() for part in _cell_text(value).replace('\\', '/').split('/') if part.strip())


def _organization_path(organization, organizations, cache, visiting=None):
    """递归构建组织完整路径，并在异常循环层级时返回空值。"""
    if organization.id in cache:
        return cache[organization.id]
    visiting = set(visiting or set())
    if organization.id in visiting:
        return ''
    visiting.add(organization.id)
    parent = organizations.get(organization.parent_id)
    parent_path = _organization_path(parent, organizations, cache, visiting) if parent else ''
    path = '/'.join(item for item in (parent_path, organization.name.strip()) if item)
    cache[organization.id] = path
    return path


def visible_organization_paths(user):
    """返回当前操作者组织范围内可用于导入的启用普通部门路径映射。"""
    scope_ids = organization_scope_ids(user)
    queryset = models.Organization.objects.filter(is_active=True).select_related('parent').order_by('id')
    if scope_ids is not None:
        queryset = queryset.filter(id__in=scope_ids) if scope_ids else queryset.none()
    organizations = {item.id: item for item in queryset}
    cache = {}
    paths = {}
    duplicates = set()
    for organization in organizations.values():
        if organization.is_default or organization.parent_id is None:
            continue
        path = _organization_path(organization, organizations, cache)
        normalized = _normalized_department_path(path)
        if not normalized:
            continue
        if normalized in paths:
            duplicates.add(normalized)
        else:
            paths[normalized] = organization
    for path in duplicates:
        paths.pop(path, None)
    return paths, duplicates


def build_user_import_template(user):
    """按当前操作者可见部门和身份生成不落盘的用户导入 Excel 模板。"""
    department_paths, _duplicates = visible_organization_paths(user)
    roles = list(models.Role.objects.filter(is_active=True).order_by('-rank', 'name'))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '用户导入'
    dictionary = workbook.create_sheet('数据字典')
    instructions = workbook.create_sheet('填写说明')

    header_fill = PatternFill('solid', fgColor='2563EB')
    required_fill = PatternFill('solid', fgColor='1D4ED8')
    for column_index, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = required_fill if header in {'用户名', '所属部门路径', '初始密码'} else header_fill
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:J{MAX_IMPORT_ROWS + 1}'
    widths = [18, 16, 38, 28, 18, 18, 18, 14, 16, 24]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width

    for row_index, path in enumerate(sorted(department_paths), start=2):
        dictionary.cell(row=row_index, column=1, value=path)
    for row_index, role in enumerate(roles, start=2):
        dictionary.cell(row=row_index, column=2, value=role.name)
        dictionary.cell(row=row_index, column=3, value=role.code)
    dictionary.cell(row=1, column=1, value='所属部门路径')
    dictionary.cell(row=1, column=2, value='身份名称')
    dictionary.cell(row=1, column=3, value='身份代码')
    dictionary.sheet_state = 'hidden'

    validations = [
        (3, sorted(department_paths), 'DepartmentPaths', 'A'),
        (6, [role.name for role in roles], 'RoleNames', 'B'),
        (8, ['是', '否'], '', ''),
        (9, ['跟随平台', '强制启用', '免于认证'], '', ''),
    ]
    for column_index, values, range_name, source_column in validations:
        validation = DataValidation(type='list', formula1=f'"{",".join(values)}"', allow_blank=column_index in {6, 8, 9})
        if range_name and values:
            workbook.defined_names.add(DefinedName(
                range_name,
                attr_text=f"'数据字典'!${source_column}$2:${source_column}${len(values) + 1}",
            ))
            validation.formula1 = f'={range_name}'
        sheet.add_data_validation(validation)
        validation.add(f'{sheet.cell(row=1, column=column_index).column_letter}2:{sheet.cell(row=1, column=column_index).column_letter}{MAX_IMPORT_ROWS + 1}')

    notes = [
        ('必填字段', '用户名、所属部门路径、初始密码。'),
        ('所属部门', '必须从下拉列表选择完整路径，系统不会自动创建不存在的部门。'),
        ('身份', '可选择身份名称；留空时使用系统默认身份。'),
        ('启用状态', '填写“是”或“否”，留空默认启用。'),
        ('OTP策略', '可填写“跟随平台、强制启用、免于认证”，留空默认跟随平台。'),
        ('安全提示', '模板包含初始密码，完成导入后请立即从本地安全删除文件。'),
        ('导入规则', f'单次最多 {MAX_IMPORT_ROWS} 行；预检查存在任何异常时整批禁止导入。'),
    ]
    for row_index, (title, content) in enumerate(notes, start=1):
        instructions.cell(row=row_index, column=1, value=title).font = Font(bold=True)
        instructions.cell(row=row_index, column=2, value=content)
    instructions.column_dimensions['A'].width = 18
    instructions.column_dimensions['B'].width = 86

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _read_upload_bytes(upload):
    """读取并限制上传文件及 ZIP 解压尺寸，拒绝伪装格式和压缩炸弹。"""
    if not upload:
        raise UserImportError('请选择要上传的 Excel 文件')
    if not str(getattr(upload, 'name', '') or '').lower().endswith('.xlsx'):
        raise UserImportError('用户导入只支持 .xlsx 文件')
    content = upload.read(MAX_IMPORT_FILE_SIZE + 1)
    if len(content) > MAX_IMPORT_FILE_SIZE:
        raise UserImportError('Excel 文件不能超过 5 MB')
    if not content:
        raise UserImportError('Excel 文件不能为空')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 200:
                raise UserImportError('Excel 文件内部条目过多')
            if sum(item.file_size for item in archive.infolist()) > MAX_IMPORT_ARCHIVE_SIZE:
                raise UserImportError('Excel 解压后内容过大')
    except BadZipFile as exc:
        raise UserImportError('Excel 文件格式无效') from exc
    return content


def _flatten_errors(detail):
    """把 DRF 字段错误转换为适合逐行预览的中文字符串列表。"""
    messages = []
    if isinstance(detail, dict):
        for field, value in detail.items():
            for message in _flatten_errors(value):
                messages.append(f'{field}：{message}')
    elif isinstance(detail, (list, tuple)):
        for value in detail:
            messages.extend(_flatten_errors(value))
    else:
        messages.append(str(detail))
    return messages


def _parse_active(value):
    """把 Excel 中的启用状态转换为布尔值，空值默认启用。"""
    text = _cell_text(value).casefold()
    if not text:
        return True, ''
    if text not in ACTIVE_VALUES:
        return True, '启用状态只能填写“是”或“否”'
    return ACTIVE_VALUES[text], ''


def _parse_otp(value):
    """把 Excel 中的 OTP 中文选项转换为用户档案策略代码。"""
    text = _cell_text(value)
    normalized = text.casefold() if text.isascii() else text
    if normalized not in OTP_VALUES:
        return 'inherit', 'OTP策略只能填写“跟随平台、强制启用、免于认证”'
    return OTP_VALUES[normalized], ''


def validate_user_import(upload, user):
    """解析并逐行预检查用户 Excel，不创建用户或保存上传文件。"""
    content = _read_upload_bytes(upload)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise UserImportError('Excel 文件无法读取，请重新下载模板填写') from exc
    if '用户导入' not in workbook.sheetnames:
        raise UserImportError('Excel 缺少“用户导入”工作表')
    sheet = workbook['用户导入']
    headers = tuple(_cell_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    if headers[:len(TEMPLATE_HEADERS)] != TEMPLATE_HEADERS:
        raise UserImportError('Excel 表头不匹配，请重新下载最新模板')

    department_paths, duplicate_paths = visible_organization_paths(user)
    roles = list(models.Role.objects.filter(is_active=True))
    roles_by_value = {}
    for role in roles:
        roles_by_value[role.code.casefold()] = role
        roles_by_value[role.name.casefold()] = role
    default_role = roles_by_value.get(serializers.default_role_code().casefold())
    seen_usernames = set()
    seen_phone_hashes = set()
    rows = []
    validated_rows = []

    for excel_row, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(TEMPLATE_HEADERS)), start=2):
        values = [cell.value for cell in cells]
        if not any(_cell_text(value) for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise UserImportError(f'单次最多导入 {MAX_IMPORT_ROWS} 行用户数据')
        errors = []
        if any(getattr(cell, 'data_type', '') == 'f' or _cell_text(cell.value).startswith('=') for cell in cells):
            errors.append('不允许使用 Excel 公式')
        username = _cell_text(values[0])
        full_name = _cell_text(values[1])
        department_path = _normalized_department_path(values[2])
        email = _cell_text(values[3])
        phone = _cell_text(values[4])
        role_value = _cell_text(values[5])
        title = _cell_text(values[6])
        is_active, active_error = _parse_active(values[7])
        otp_policy, otp_error = _parse_otp(values[8])
        password = _cell_text(values[9])
        if active_error:
            errors.append(active_error)
        if otp_error:
            errors.append(otp_error)
        if not username:
            errors.append('用户名不能为空')
        username_key = username.casefold()
        if username_key in seen_usernames:
            errors.append('Excel 内用户名重复')
        seen_usernames.add(username_key)
        if phone:
            try:
                normalized_phone = normalize_mobile_phone(phone)
                current_phone_hash = phone_lookup_hash(normalized_phone)
                if current_phone_hash in seen_phone_hashes:
                    errors.append('Excel 内手机号重复')
                seen_phone_hashes.add(current_phone_hash)
                phone = normalized_phone
            except ValueError as exc:
                errors.append(str(exc))
        if not department_path:
            errors.append('所属部门路径不能为空')
        if department_path in duplicate_paths:
            errors.append('所属部门路径不唯一，请联系管理员调整同名部门')
        organization = department_paths.get(department_path)
        if department_path and not organization and department_path not in duplicate_paths:
            errors.append('所属部门不存在、已停用或不在当前管理范围')
        role = roles_by_value.get(role_value.casefold()) if role_value else default_role
        if not role:
            errors.append('身份不存在、已停用或系统未配置默认身份')
        if not password:
            errors.append('初始密码不能为空')

        payload = {
            'username': username,
            'first_name': '',
            'last_name': full_name,
            'email': email,
            'phone': phone,
            'role': role.code if role else role_value,
            'title': title,
            'is_active': is_active,
            'otp_policy': otp_policy,
            'password': password,
            'department_id': organization.id if organization else None,
        }
        serializer = serializers.UserSerializer(data=payload)
        if not serializer.is_valid():
            errors.extend(_flatten_errors(serializer.errors))
        row = {
            'row_number': excel_row,
            'username': username or '—',
            'name': full_name or '—',
            'department': department_path or '—',
            'status': 'invalid' if errors else 'valid',
            'errors': list(dict.fromkeys(errors)),
        }
        rows.append(row)
        if not errors:
            validated_rows.append(payload)

    if not rows:
        raise UserImportError('Excel 中没有可导入的用户数据')
    invalid_rows = sum(1 for row in rows if row['errors'])
    active_rows = sum(1 for payload in validated_rows if payload['is_active'])
    if not invalid_rows:
        denial = license_user_limit_denial(active_rows)
        if denial:
            message = denial['detail']
            for row, payload in zip(rows, validated_rows):
                if payload['is_active']:
                    row['status'] = 'invalid'
                    row['errors'] = [message]
            invalid_rows = active_rows or len(rows)
            validated_rows = []
    return {
        'total_rows': len(rows),
        'valid_rows': len(rows) - invalid_rows,
        'invalid_rows': invalid_rows,
        'can_import': invalid_rows == 0,
        'rows': rows,
        '_validated_rows': validated_rows,
    }


def public_validation_result(result):
    """移除内部标准化载荷，只向前端返回安全的汇总和逐行错误。"""
    return {key: value for key, value in result.items() if not key.startswith('_')}


def import_users(upload, user):
    """重新校验同一 Excel，并在单个事务中整批创建用户及加密档案。"""
    result = validate_user_import(upload, user)
    if not result['can_import']:
        raise UserImportError('预检查发现异常，暂不允许导入', public_validation_result(result))
    payloads = result['_validated_rows']
    try:
        with transaction.atomic():
            models.SystemSetting.objects.select_for_update().filter(key='license.management').first()
            denial = license_user_limit_denial(sum(1 for payload in payloads if payload['is_active']))
            if denial:
                raise UserImportError(denial['detail'])
            created_users = []
            for payload in payloads:
                serializer = serializers.UserSerializer(data=payload)
                serializer.is_valid(raise_exception=True)
                created_users.append(serializer.save())
    except (IntegrityError, DRFValidationError) as exc:
        raise UserImportError('导入期间用户数据发生变化，请重新预检查') from exc
    department_count = len({user.profile.organization_id for user in created_users})
    return {
        'detail': f'成功导入 {len(created_users)} 个用户',
        'created_count': len(created_users),
        'department_count': department_count,
    }
