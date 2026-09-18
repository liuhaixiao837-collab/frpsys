from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from ops.models import AuditLog, Organization, Role, UserProfile
from ops.services.user_import import TEMPLATE_HEADERS


class UserImportTests(TestCase):
    """验证用户模板下载、预检查、跨部门导入和整批回滚。"""

    def setUp(self):
        """创建管理员、两个部门和可分配身份。"""
        self.company = Organization.objects.create(
            name='测试公司', slug='import-company', org_type='company', is_default=True,
        )
        self.development = Organization.objects.create(
            name='研发部', slug='import-development', org_type='department', parent=self.company,
        )
        self.operations = Organization.objects.create(
            name='运维部', slug='import-operations', org_type='department', parent=self.company,
        )
        Role.objects.update_or_create(
            code='member',
            defaults={'name': '普通用户', 'rank': 10, 'is_active': True},
        )
        self.admin = User.objects.create_superuser(
            username='admin', email='admin@example.com', password='StrongPassword123!',
        )
        UserProfile.objects.create(user=self.admin, organization=self.company, role='admin')
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

    def workbook_upload(self, rows, name='users.xlsx'):
        """把测试行写入内存 XLSX 并包装为 multipart 上传文件。"""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '用户导入'
        sheet.append(TEMPLATE_HEADERS)
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            name,
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def valid_rows(self):
        """返回分别指向研发部和运维部的两行合法用户数据。"""
        return [
            ['developer', '开发用户', '测试公司/研发部', 'dev@example.com', '13800138001', '普通用户', '开发工程师', '是', '跟随平台', 'StrongPassword123!'],
            ['operator', '运维用户', '测试公司/运维部', 'ops@example.com', '13800138002', 'member', '运维工程师', '否', '免于认证', 'StrongPassword123!'],
        ]

    def test_template_is_generated_with_department_dictionary(self):
        """模板必须直接下载且包含当前部门路径和填写说明。"""
        response = self.client.get('/api/v1/users/import/template/')

        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment;', response['Content-Disposition'])
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook['用户导入']['A1'].value, '用户名')
        self.assertIn('数据字典', workbook.sheetnames)
        department_values = {
            workbook['数据字典'].cell(row=row, column=1).value
            for row in range(2, workbook['数据字典'].max_row + 1)
        }
        self.assertIn('测试公司/研发部', department_values)
        self.assertIn('测试公司/运维部', department_values)

    def test_precheck_reports_invalid_department_without_creating_users(self):
        """部门路径错误时预检查应返回逐行异常且不得创建任何用户。"""
        upload = self.workbook_upload([
            ['invalid-user', '错误用户', '测试公司/不存在部门', '', '', 'member', '', '是', '跟随平台', 'StrongPassword123!'],
        ])

        response = self.client.post('/api/v1/users/import/precheck/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertFalse(payload['can_import'])
        self.assertEqual(payload['invalid_rows'], 1)
        self.assertIn('所属部门不存在', '；'.join(payload['rows'][0]['errors']))
        self.assertFalse(User.objects.filter(username='invalid-user').exists())

    def test_confirm_imports_users_into_excel_departments(self):
        """全部预检查通过后应按 Excel 部门路径整批创建用户和加密档案。"""
        precheck = self.client.post(
            '/api/v1/users/import/precheck/',
            {'file': self.workbook_upload(self.valid_rows())},
            format='multipart',
        )
        self.assertEqual(precheck.status_code, 200, precheck.json())
        self.assertTrue(precheck.json()['can_import'])

        response = self.client.post(
            '/api/v1/users/import/confirm/',
            {'file': self.workbook_upload(self.valid_rows())},
            format='multipart',
        )

        self.assertEqual(response.status_code, 201, response.json())
        self.assertEqual(response.json()['created_count'], 2)
        self.assertEqual(UserProfile.objects.get(user__username='developer').organization, self.development)
        operator = User.objects.get(username='operator')
        self.assertEqual(operator.profile.organization, self.operations)
        self.assertFalse(operator.is_active)
        self.assertEqual(operator.profile.otp_policy, 'exempt')
        audit = AuditLog.objects.filter(action='User.import').latest('id')
        self.assertEqual(audit.detail['result'], 'success')
        self.assertNotIn('password', str(audit.detail).lower())

    def test_confirm_rejects_entire_file_when_any_row_is_invalid(self):
        """混合合法和异常行时确认导入必须整批拒绝并保持数据库不变。"""
        rows = self.valid_rows()
        rows[1][2] = '测试公司/不存在部门'

        response = self.client.post(
            '/api/v1/users/import/confirm/',
            {'file': self.workbook_upload(rows)},
            format='multipart',
        )

        self.assertEqual(response.status_code, 400, response.json())
        self.assertFalse(User.objects.filter(username__in=['developer', 'operator']).exists())
        audit = AuditLog.objects.filter(action='User.import').latest('id')
        self.assertEqual(audit.detail, {'result': 'failed'})


class UserImportPermissionTests(TestCase):
    """验证普通用户必须获得用户导入操作权限才能访问全部导入接口。"""

    def setUp(self):
        """创建不具备导入权限的普通用户。"""
        organization = Organization.objects.create(
            name='权限公司', slug='import-permission-company', org_type='company', is_default=True,
        )
        self.user = User.objects.create_user(username='manager', password='StrongPassword123!')
        UserProfile.objects.create(user=self.user, organization=organization, role='member')
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_import_template_requires_import_action_permission(self):
        """仅登录但未授予导入操作权限时模板接口必须拒绝访问。"""
        response = self.client.get('/api/v1/users/import/template/')

        self.assertEqual(response.status_code, 403)
