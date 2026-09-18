import json
from bisect import bisect_left
from datetime import timedelta
from io import BytesIO

from django.conf import settings
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError


MAX_LOG_EXPORT_DAYS = 180

ACTION_LABELS = {
    'auth.login': '登录系统',
    'auth.login.success': '登录成功',
    'auth.login.failed': '登录失败',
    'auth.switch_org': '切换部门',
    'User.create': '新增用户',
    'User.update': '编辑用户',
    'User.delete': '删除用户',
    'User.import': '批量导入用户',
    'User.otp_bind': '绑定 OTP',
    'User.otp_reset': '重置 OTP',
    'Organization.create': '新增部门',
    'Organization.update': '编辑部门',
    'Organization.delete': '删除部门',
    'Organization.add_member': '添加部门成员',
    'Organization.invite_member': '邀请部门成员',
    'Organization.remove_member': '移出部门成员',
    'PermissionPolicy.create': '新增权限策略',
    'PermissionPolicy.update': '编辑权限策略',
    'PermissionPolicy.delete': '删除权限策略',
    'SystemSetting.create': '新增平台设置',
    'SystemSetting.update': '编辑平台设置',
    'SystemSetting.delete': '删除平台设置',
    'SystemSetting.license_import': '导入 Licence',
    'SystemSetting.test_email': '发送测试邮件',
    'SystemSetting.test_sms': '发送测试短信',
    'SystemSetting.test_llm': '检测 LLM 服务',
    'MenuOrder.update': '调整主菜单顺序',
    'AuditLog.export': '导出日志',
}

RESOURCE_NAMES = {
    'User': '用户',
    'Organization': '部门',
    'PermissionPolicy': '权限策略',
    'SystemSetting': '平台设置',
    'MenuOrder': '菜单排序',
    'AuditLog': '日志',
    'auth': '账号安全',
}

OPERATION_WORDS = {
    'create': '新增资源',
    'update': '编辑资源',
    'delete': '删除资源',
    'enable': '启用资源',
    'disable': '禁用资源',
    'approve': '通过审批',
    'reject': '驳回审批',
    'cancel': '取消操作',
    'connect': '建立连接',
    'disconnect': '断开连接',
    'execute': '执行操作',
    'download': '下载文件',
    'upload': '上传文件',
    'cleanup': '清理数据',
}

LOGIN_METHOD_LABELS = {
    'password': '账号密码',
    'password+otp': '密码 + OTP',
    'sso': '单点登录',
    'ldap': 'LDAP',
    'oauth': 'OAuth',
    'sms': '短信验证码',
}


def parse_log_time_range(params, require_complete=False, enforce_limit=False):
    """解析日志开始和结束时间，并按导出场景校验 180 天限制。"""
    start_text = str(params.get('start_time') or '').strip()
    end_text = str(params.get('end_time') or '').strip()
    if require_complete and (not start_text or not end_text):
        raise ValidationError('导出前请选择完整的开始时间和结束时间')
    if bool(start_text) != bool(end_text):
        raise ValidationError('开始时间和结束时间必须同时填写')
    if not start_text:
        return None, None
    start_time = parse_datetime(start_text)
    end_time = parse_datetime(end_text)
    if not start_time or not end_time:
        raise ValidationError('时间格式无效，请重新选择时间范围')
    if timezone.is_naive(start_time):
        start_time = timezone.make_aware(start_time, timezone.get_current_timezone())
    if timezone.is_naive(end_time):
        end_time = timezone.make_aware(end_time, timezone.get_current_timezone())
    if start_time >= end_time:
        raise ValidationError('结束时间必须晚于开始时间')
    if enforce_limit and end_time - start_time > timedelta(days=MAX_LOG_EXPORT_DAYS):
        raise ValidationError(f'单次最多导出 {MAX_LOG_EXPORT_DAYS} 天日志，请缩小时间范围')
    return start_time, end_time


def apply_log_time_range(queryset, start_time, end_time):
    """把已校验的时间范围应用到日志查询集。"""
    if start_time:
        queryset = queryset.filter(created_at__gte=start_time)
    if end_time:
        queryset = queryset.filter(created_at__lte=end_time)
    return queryset


def audit_action_label(action):
    """将审计动作代码转换为适合报表阅读的中文名称。"""
    if action in ACTION_LABELS:
        return ACTION_LABELS[action]
    if str(action).lower().startswith('http.'):
        method = str(action).rsplit('.', 1)[-1].lower()
        return {'post': '新增或提交操作', 'put': '编辑操作', 'patch': '编辑操作', 'delete': '删除操作'}.get(
            method,
            '执行业务操作',
        )
    operation = str(action).rsplit('.', 1)[-1]
    operation_label = OPERATION_WORDS.get(operation)
    if not operation_label:
        return str(action or '执行业务操作')
    resource_name = next((name for code, name in RESOURCE_NAMES.items() if code in str(action).split('.')), '业务资源')
    return operation_label.replace('资源', resource_name)


def audit_resource_label(log):
    """将审计资源代码转换为包含类型和名称的业务对象说明。"""
    action_parts = str(log.action or '').replace('bastion.', '').split('.')
    resource_type = next((RESOURCE_NAMES[part] for part in action_parts if part in RESOURCE_NAMES), '其他资源')
    resource = str(log.resource or '').strip()
    if not resource:
        return resource_type
    if resource.startswith('/api/'):
        return resource_type
    if resource == resource_type or resource.startswith(f'{resource_type}：'):
        return resource
    return f'{resource_type}：{resource}'


def safe_detail(detail):
    """递归移除凭据类字段并返回适合写入 Excel 的日志详情。"""
    if isinstance(detail, dict):
        return {
            str(key): safe_detail(value)
            for key, value in detail.items()
            if not any(word in str(key).lower() for word in ('password', 'secret', 'token', 'credential'))
        }
    if isinstance(detail, list):
        return [safe_detail(value) for value in detail]
    return detail


def detail_text(detail):
    """将安全日志详情序列化为紧凑、可检索的中文文本。"""
    value = safe_detail(detail or {})
    if not value:
        return '无补充信息'
    return json.dumps(value, ensure_ascii=False, separators=('，', '：'))


def login_status(log):
    """返回登录日志的成功或失败状态。"""
    status_value = str((log.detail or {}).get('status') or '').lower()
    return 'failed' if status_value == 'failed' or str(log.action).endswith('.failed') else 'success'


def login_method(log):
    """返回登录日志中的认证方式代码。"""
    return str((log.detail or {}).get('auth_method') or 'password').lower()


def login_reason(log):
    """返回登录日志中的结果说明。"""
    default_reason = '登录成功' if login_status(log) == 'success' else '登录失败'
    return str((log.detail or {}).get('reason') or default_reason)


def remove_paired_technical_logs(rows):
    """移除与中文业务日志成对产生的 HTTP 兜底日志，保持页面和导出口径一致。"""
    business_times = {}
    for log in rows:
        if not str(log.action or '').lower().startswith('http.'):
            business_times.setdefault(log.actor, []).append(log.created_at.timestamp())
    for timestamps in business_times.values():
        timestamps.sort()
    result = []
    for log in rows:
        if not str(log.action or '').lower().startswith('http.'):
            result.append(log)
            continue
        timestamps = business_times.get(log.actor, [])
        occurred_at = log.created_at.timestamp()
        index = bisect_left(timestamps, occurred_at - 3)
        if index < len(timestamps) and abs(timestamps[index] - occurred_at) <= 3:
            continue
        result.append(log)
    return result


def filter_export_logs(rows, kind, params):
    """按照页面当前关键词、状态和登录方式筛选待导出的全部日志。"""
    result = remove_paired_technical_logs(rows) if kind == 'system' else list(rows)
    query = str(params.get('q') or '').strip().lower()
    status_filter = str(params.get('status') or '').strip().lower()
    method_filter = str(params.get('method') or '').strip().lower()
    filtered = []
    for log in result:
        if kind == 'user':
            status_value = login_status(log)
            method_value = login_method(log)
            if status_filter and status_value != status_filter:
                continue
            if method_filter and method_value != method_filter:
                continue
            values = [
                log.actor,
                log.ip_address,
                '成功' if status_value == 'success' else '失败',
                LOGIN_METHOD_LABELS.get(method_value, method_value),
                login_reason(log),
            ]
        else:
            values = [
                log.actor,
                log.ip_address,
                audit_action_label(log.action),
                audit_resource_label(log),
                detail_text(log.detail),
            ]
        if query and not any(query in str(value or '').lower() for value in values):
            continue
        filtered.append(log)
    return filtered


def local_excel_time(value):
    """将 Django 时间转换为 Excel 可识别的本地无时区时间。"""
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.replace(tzinfo=None)


def export_filter_description(kind, params):
    """生成写入工作簿顶部的当前筛选条件说明。"""
    filters = []
    query = str(params.get('q') or '').strip()
    if query:
        filters.append(f'关键词：{query}')
    if kind == 'user':
        status_value = str(params.get('status') or '').strip()
        method_value = str(params.get('method') or '').strip()
        if status_value:
            filters.append(f"登录状态：{'成功' if status_value == 'success' else '失败'}")
        if method_value:
            filters.append(f'登录方式：{LOGIN_METHOD_LABELS.get(method_value, method_value)}')
    return '；'.join(filters) if filters else '其他筛选：无'


def build_log_workbook(rows, kind, start_time, end_time, params):
    """生成带品牌化标题、筛选摘要、冻结表头和打印设置的日志工作簿。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '用户日志' if kind == 'user' else '系统日志'
    sheet.sheet_view.showGridLines = False
    platform_name = getattr(settings, 'PLATFORM_DISPLAY_NAME', '平台')
    title = f"{platform_name} · {sheet.title}导出"
    headers = (
        ['序号', '登录时间', '登录用户', '登录状态', '登录方式', '来源地址', '结果说明']
        if kind == 'user'
        else ['序号', '操作时间', '操作者', '操作类型', '操作对象', '来源地址', '操作详情']
    )
    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f'A1:{last_column}1')
    sheet['A1'] = title
    sheet['A1'].font = Font(name='微软雅黑', size=18, bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='17365D')
    sheet['A1'].alignment = Alignment(horizontal='left', vertical='center')
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells(f'A2:{last_column}2')
    sheet['A2'] = (
        f"统计范围：{timezone.localtime(start_time).strftime('%Y-%m-%d %H:%M:%S')} 至 "
        f"{timezone.localtime(end_time).strftime('%Y-%m-%d %H:%M:%S')}　｜　共 {len(rows)} 条"
    )
    sheet['A2'].font = Font(name='微软雅黑', size=10, color='334155')
    sheet['A2'].fill = PatternFill('solid', fgColor='EAF2F8')
    sheet['A2'].alignment = Alignment(horizontal='left', vertical='center')
    sheet.row_dimensions[2].height = 24

    sheet.merge_cells(f'A3:{last_column}3')
    sheet['A3'] = f"{export_filter_description(kind, params)}　｜　导出时间：{timezone.localtime().strftime('%Y-%m-%d %H:%M:%S')}"
    sheet['A3'].font = Font(name='微软雅黑', size=9, color='64748B')
    sheet['A3'].alignment = Alignment(horizontal='left', vertical='center')
    sheet.row_dimensions[3].height = 22

    header_fill = PatternFill('solid', fgColor='2563EB')
    header_border = Border(bottom=Side(style='medium', color='1D4ED8'))
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column_index, value=header)
        cell.font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[5].height = 26

    thin_border = Border(bottom=Side(style='thin', color='DCE6F1'))
    for index, log in enumerate(rows, start=1):
        status_value = login_status(log) if kind == 'user' else ''
        method_value = login_method(log) if kind == 'user' else ''
        values = (
            [
                index,
                local_excel_time(log.created_at),
                log.actor or '未知用户',
                '成功' if status_value == 'success' else '失败',
                LOGIN_METHOD_LABELS.get(method_value, method_value or '其他方式'),
                log.ip_address or '未知地址',
                login_reason(log),
            ]
            if kind == 'user'
            else [
                index,
                local_excel_time(log.created_at),
                log.actor or '系统服务',
                audit_action_label(log.action),
                audit_resource_label(log),
                log.ip_address or '未知地址',
                detail_text(log.detail),
            ]
        )
        row_number = index + 5
        row_fill = PatternFill('solid', fgColor='F7FAFC' if index % 2 == 0 else 'FFFFFF')
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column_index, value=value)
            cell.font = Font(name='微软雅黑', size=9, color='1E293B')
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal='center' if column_index in {1, 2, 4, 6} else 'left',
                vertical='center',
                wrap_text=column_index == len(headers),
            )
        sheet.cell(row=row_number, column=2).number_format = 'yyyy-mm-dd hh:mm:ss'
        if kind == 'user':
            status_cell = sheet.cell(row=row_number, column=4)
            status_cell.font = Font(
                name='微软雅黑',
                size=9,
                bold=True,
                color='047857' if status_value == 'success' else 'B91C1C',
            )
        sheet.row_dimensions[row_number].height = 24

    widths = [8, 22, 20, 18, 42, 18, 58]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.freeze_panes = 'A6'
    sheet.auto_filter.ref = f'A5:{last_column}{max(5, len(rows) + 5)}'
    sheet.print_title_rows = '1:5'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = f'{platform_name}日志报表 · 第 &P / &N 页'
    sheet.oddFooter.center.size = 9
    workbook.properties.creator = platform_name
    workbook.properties.title = title
    workbook.properties.subject = '日志安全审计导出'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
