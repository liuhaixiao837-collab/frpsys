from datetime import timedelta
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from ops.models import AuditLog, Organization, UserProfile


class LogExportTests(TestCase):
    """验证用户日志和系统日志的时间筛选及 Excel 导出。"""

    def setUp(self):
        """创建超级管理员和不同类型、不同时间的日志数据。"""
        root = Organization.objects.filter(is_default=True).first()
        if not root:
            root = Organization.objects.create(
                name='默认组织', slug='log-export-default', org_type='company', is_default=True,
            )
        self.admin = User.objects.create_superuser(
            username='admin', email='admin@example.com', password='StrongPassword123!',
        )
        UserProfile.objects.create(user=self.admin, organization=root, role='admin')
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        self.now = timezone.now()
        self.user_success = AuditLog.objects.create(
            organization=root,
            actor='alice',
            action='auth.login.success',
            resource='alice',
            ip_address='192.0.2.10',
            detail={'status': 'success', 'auth_method': 'password+otp', 'reason': '登录成功'},
            created_at=self.now - timedelta(days=2),
        )
        self.user_failed = AuditLog.objects.create(
            organization=root,
            actor='bob',
            action='auth.login.failed',
            resource='bob',
            ip_address='192.0.2.11',
            detail={'status': 'failed', 'auth_method': 'password', 'reason': '密码错误'},
            created_at=self.now - timedelta(days=1),
        )
        self.system_log = AuditLog.objects.create(
            organization=root,
            actor='admin',
            action='User.update',
            resource='alice',
            ip_address='192.0.2.12',
            detail={'user_id': 9, 'status': 'enabled'},
            created_at=self.now - timedelta(hours=12),
        )

    def export_range(self, days=7):
        """返回覆盖当前测试日志的标准导出时间参数。"""
        return {
            'start_time': (self.now - timedelta(days=days)).isoformat(),
            'end_time': (self.now + timedelta(minutes=1)).isoformat(),
        }

    def test_user_log_list_filters_by_time(self):
        """用户日志列表应只返回开始和结束时间内的数据。"""
        response = self.client.get('/api/v1/user-logs/', {
            'start_time': (self.now - timedelta(days=1, hours=12)).isoformat(),
            'end_time': (self.now + timedelta(minutes=1)).isoformat(),
        })

        self.assertEqual(response.status_code, 200, response.json())
        rows = response.json()['results']
        self.assertEqual([row['actor'] for row in rows], ['bob'])

    def test_user_log_export_uses_current_status_filter(self):
        """用户日志 Excel 应导出完整范围内符合当前状态筛选的全部数据。"""
        params = {**self.export_range(), 'status': 'failed'}
        response = self.client.get('/api/v1/user-logs/export/', params)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook['用户日志']
        self.assertEqual(sheet['A1'].value, '基础平台 · 用户日志导出')
        self.assertEqual(sheet['D5'].value, '登录状态')
        self.assertEqual(sheet['C6'].value, 'bob')
        self.assertEqual(sheet['D6'].value, '失败')
        self.assertEqual(sheet.max_row, 6)

    def test_system_log_export_creates_styled_workbook(self):
        """系统日志 Excel 应包含筛选摘要、中文业务字段和冻结表头。"""
        response = self.client.get('/api/v1/system-logs/export/', self.export_range())

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook['系统日志']
        self.assertEqual(sheet['A1'].value, '基础平台 · 系统日志导出')
        self.assertIn('共 1 条', sheet['A2'].value)
        self.assertEqual(sheet['D6'].value, '编辑用户')
        self.assertEqual(sheet.freeze_panes, 'A6')
        self.assertEqual(sheet.auto_filter.ref, 'A5:G6')

    def test_export_rejects_range_longer_than_180_days(self):
        """任何日志导出接口都必须拒绝超过一百八十天的时间范围。"""
        response = self.client.get('/api/v1/system-logs/export/', {
            'start_time': (self.now - timedelta(days=181)).isoformat(),
            'end_time': self.now.isoformat(),
        })

        self.assertEqual(response.status_code, 400, response.json())
        self.assertIn('180', ' '.join(response.json()['detail']))
